import math
import struct
import sys
import wave
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from mooc_m1.adapters import LocalModelAdapter
from mooc_m1.core import Failure, digest, run, write_json
from mooc_m1.media import MediaTools
from mooc_m1.ppt import parse_pptx
from mooc_m1.preflight import preflight, check_constraints


def wav(path, silent=False, seconds=2):
    with wave.open(str(path), "wb") as f:
        f.setnchannels(1)
        f.setsampwidth(2)
        f.setframerate(16000)
        f.writeframes(b"".join(struct.pack("<h", 0 if silent else round(4000 * math.sin(2 * math.pi * 440 * i / 16000))) for i in range(int(seconds * 16000))))


def test_ppt_preserves_order_notes_images_and_stable_ids(tmp_path):
    prs = Presentation()
    for i in range(10):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if i != 9:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = f"原页 {i+1}：数组 [1,2]，公式 (a+b)"
            slide.notes_slide.notes_text_frame.text = f"原备注 {i+1} 不润色"
        else:
            image = tmp_path / "image.png"
            Image.new("RGB", (100, 100), "blue").save(image)
            slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    path = tmp_path / "slides.pptx"
    prs.save(path)
    a = parse_pptx(path, tmp_path / "a")
    b = parse_pptx(path, tmp_path / "b")
    assert a["page_count"] == 10 and not a["rendered"]
    assert [p["source_index"] for p in a["pages"]] == list(range(1, 11))
    assert a["pages"][0]["notes_original"] == "原备注 1 不润色"
    assert "[1,2]" in a["pages"][0]["body"] and "(a+b)" in a["pages"][0]["body"]
    assert [p["source_page_id"] for p in a["pages"]] == [p["source_page_id"] for p in b["pages"]]
    assert a["pages"][-1]["body"] == "" and a["pages"][-1]["images"]
    assert "IMAGE_ONLY_PAGE" in {i["code"] for i in a["pages"][-1]["anomalies"]}


def test_fake_ppt_extension_rejected(tmp_path):
    path = tmp_path / "bad.pptx"
    path.write_bytes(b"not a ppt")
    with pytest.raises(Failure, match="Invalid PPTX"):
        parse_pptx(path, tmp_path / "out")


def test_ppt_limit_is_configurable_and_cli_failure_preserves_directory(tmp_path):
    from mooc_m1.core import read_json
    prs = Presentation()
    for i in range(51):
        prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "long.pptx"
    prs.save(path)
    with pytest.raises(Failure) as error:
        parse_pptx(path, tmp_path / "default")
    assert error.value.code == "INPUT_INCOMPATIBLE"
    assert parse_pptx(path, tmp_path / "configured", max_pages=63)["page_count"] == 51
    output = tmp_path / "cli-output"
    with pytest.raises(Failure):
        run([sys.executable, "-m", "mooc_m1", "ppt", str(path), "--output", str(output)])
    assert output.is_dir()
    assert read_json(output / "error.json")["error"]["code"] == "INPUT_INCOMPATIBLE"


@pytest.mark.parametrize("silent,code", [(True, "ABNORMAL_SILENCE"), (False, None)])
def test_actual_ffmpeg_silence_and_duration(tmp_path, media, silent, code):
    path = tmp_path / "audio.wav"
    wav(path, silent)
    if code:
        with pytest.raises(Failure) as error:
            media.inspect(path, "audio")
        assert error.value.code == code
    else:
        result = media.inspect(path, "audio")
        assert result["duration_seconds"] == pytest.approx(2)
        assert result["silence_ratio"] == 0
        with pytest.raises(Failure) as error:
            media.inspect(path, "audio", expected_duration=30)
        assert error.value.code == "DURATION_MISMATCH"


@pytest.mark.parametrize("content,code", [(b"", "NO_VALID_MEDIA"), (b"invalid", "UNDECODABLE")])
def test_bad_audio(tmp_path, media, content, code):
    path = tmp_path / "audio.wav"
    path.write_bytes(content)
    with pytest.raises(Failure) as error:
        media.inspect(path, "audio")
    assert error.value.code == code


def test_photo_type_from_content_not_suffix(tmp_path, media):
    path = tmp_path / "photo.jpg"
    Image.new("RGB", (300, 400)).save(path, format="PNG")
    assert media.inspect(path, "photo")["format"] == "PNG"
    path.write_bytes(b"bad")
    with pytest.raises(Failure) as error:
        media.inspect(path, "photo")
    assert error.value.code == "UNDECODABLE"


def test_no_video_stream_does_not_become_photo_mode(tmp_path, media):
    path = tmp_path / "fake.mp4"
    wav(path)
    with pytest.raises(Failure) as error:
        media.inspect(path, "video")
    assert error.value.code == "INPUT_INCOMPATIBLE"


def test_missing_models_block_and_persist_failure(tmp_path, media):
    adapter = LocalModelAdapter("photo", None, media, tmp_path)
    result = adapter.generate({})
    assert result["state"] == "failed" and result["output"] is None
    assert result["error"]["code"] == "SERVICE_CONFIG_MISSING"
    assert not list(tmp_path.rglob("*.mp4"))
    # Recreate adapter: state is read from disk rather than an in-memory dictionary.
    assert LocalModelAdapter("photo", None, media, tmp_path).status(result["request_id"])["state"] == "failed"


def test_missing_weights(tmp_path, media):
    a = LocalModelAdapter("video", {"repo": str(tmp_path), "python": sys.executable,
        "code_revision": "unknown", "weight_version": "unverified", "usage_review": {"accepted": True},
        "weights": [{"path": str(tmp_path / "unavailable.pth"), "sha256": "0" * 64}]}, media, tmp_path)
    h = a.health()
    assert not h["ready"] and "MODEL_NOT_READY" in {e["code"] for e in h["blockers"]}


def test_preflight_no_generation_and_all_missing_inputs(tmp_path, media):
    adapters = {role: LocalModelAdapter(role, None, media, tmp_path) for role in ["tts", "photo", "video"]}
    result = preflight({}, adapters, media, tmp_path)
    assert not result["ready"] and result["generation_called"] is False
    assert len([e for e in result["blockers"] if e["code"] == "INPUT_MISSING"]) == 4
    assert result["estimated_seconds"] is None


def test_constraints_not_replaced_by_suggestion():
    with pytest.raises(Failure) as error:
        check_constraints("tts", {"reference_audio": {"duration_seconds": 25}}, {"reference_audio": {"max_seconds": 10}})
    assert error.value.code == "INPUT_INCOMPATIBLE"


def test_timeout_is_real_subprocess_fault():
    with pytest.raises(Failure) as error:
        run([sys.executable, "-c", "import time; time.sleep(10)"], timeout=0.1)
    assert error.value.code == "TIMEOUT"


def test_timeout_retains_actual_process_log(tmp_path):
    log = tmp_path / "native.log"
    with pytest.raises(Failure) as error:
        run([sys.executable, "-c", "import time; print('native step before timeout',flush=True); time.sleep(10)"],
            timeout=0.7, log_path=log)
    assert error.value.code == "TIMEOUT"
    assert "native step before timeout" in log.read_text(encoding="utf-8")


def test_process_failure_no_success_fallback():
    with pytest.raises(Failure) as error:
        run([sys.executable, "-c", "raise RuntimeError('forced model service fault')"])
    assert error.value.code == "RUN_FAILED"


def test_out_of_memory_classified():
    with pytest.raises(Failure) as error:
        run([sys.executable, "-c", "raise RuntimeError('CUDA out of memory')"])
    assert error.value.code == "RESOURCE_INSUFFICIENT"


def test_short_completely_silent_audio_is_rejected(tmp_path, media):
    path = tmp_path / "short.wav"
    wav(path, silent=True, seconds=0.2)
    with pytest.raises(Failure) as error:
        media.inspect(path, "audio")
    assert error.value.code == "ABNORMAL_SILENCE"


def test_running_record_from_dead_owner_is_failed_after_restart(tmp_path, media):
    import uuid
    request_id = str(uuid.uuid4())
    status_path = tmp_path / "requests" / request_id / "status.json"
    run([sys.executable, "-c",
         "import os,sys; from mooc_m1.core import write_json; write_json(sys.argv[1], {'request_id':sys.argv[2], 'state':'running','pid':os.getpid(),'output':None})",
         str(status_path), request_id])
    result = LocalModelAdapter("video", None, media, tmp_path).status(request_id)
    assert result["state"] == "failed"
    assert result["error"]["code"] == "RUN_FAILED"


def test_current_owner_is_alive_without_signalling_it():
    import os
    assert LocalModelAdapter._pid_alive(os.getpid())


def test_gpt_invalid_custom_yaml_blocks_fallback(tmp_path, media):
    config_path = tmp_path / "tts.yaml"
    config_path.write_text("custom:\n  version: v2\n", encoding="utf-8")
    adapter = LocalModelAdapter("tts", {"repo": str(tmp_path), "python": sys.executable,
        "code_revision": "unavailable", "weight_version": "unknown", "weights": [],
        "usage_review": {"accepted": False}, "tts_config": str(config_path)}, media, tmp_path)
    result = adapter.health()
    assert not result["ready"]
    assert any(e["message"] == "Invalid explicit GPT-SoVITS custom configuration" for e in result["blockers"])
