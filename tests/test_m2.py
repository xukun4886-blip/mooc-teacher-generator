import copy
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from mooc_m1.core import Failure
from mooc_m2.api import create_app
from mooc_m2.content import (uid, current, new_scene, add_version, speech, structure, review,
                             effective, audition_key, config_patch, import_scripts)
from mooc_m2.store import Store


@pytest.fixture
def app(tmp_path):
    config = {"storage": str(tmp_path / "storage"), "m1_config": "__missing__", "max_ppt_pages": 50,
              "limits_mb": {"pptx": 100, "photo": 20, "video": 500, "reference_audio": 50}, "ai": {}}
    return create_app(config, worker=False)


@pytest.fixture
def client(app):
    with TestClient(app) as c:
        assert c.post("/api/session", json={"token": app.state.service.store.token}).status_code == 200
        yield c


def seed(app, count=3):
    store = app.state.service.store
    p = store.create("工程测试")
    with store.edit(p["id"], p["revision"]) as p:
        p["slides"] = [{"source_page_id": f"hash:slide-{i}", "source_index": i, "body": f"原页 {i}", "notes_original": f"原备注 {i} [0,1] [1,2,3] [1]", "image": "missing.png", "anomalies": []} for i in range(1, count + 1)]
        p["scenes"] = [new_scene([s["source_page_id"]]) for s in p["slides"]]
    return store.get(p["id"])


def cmd(client, p, action, data, status=200):
    r = client.post(f"/api/projects/{p['id']}/commands", json={"revision": p["revision"], "action": action, "data": data})
    assert r.status_code == status, r.text
    return r.json()


def checks(s):
    return {"knowledge_points": ["区间与数组"], "questions": [{"question": "什么是区间？", "answer": "由边界定义的集合", "source_page_id": s["source_page_ids"][0]}], "terms": [{"text": "[0,1]", "reading": "零到一的闭区间"}], "pending": []}


def test_auth_and_cross_site(app):
    c = TestClient(app)
    assert c.get("/api/projects").status_code == 401
    assert c.post("/api/session", json={"token": "wrong"}).status_code == 401
    assert c.post("/api/session", headers={"Origin": "https://evil.invalid"}, json={"token": app.state.service.store.token}).status_code == 403
    assert c.get("/api/projects", headers={"Host": "evil.invalid"}).status_code == 403


def test_project_reopen_conflict_copy_delete(app, client):
    p = seed(app)
    p2 = cmd(client, p, "config", {"course_name": "更新课程", "audience": "本科生"})
    cmd(client, p, "config", {"course_name": "过期"}, 409)
    reopened = Store(app.state.service.store.root).get(p["id"])
    assert reopened["name"] == "更新课程" and reopened["revision"] == p2["revision"]
    target = app.state.service.store.file(p["id"], "assets/a/test.txt")
    target.parent.mkdir(parents=True)
    target.write_text("owned")
    clone = client.post(f"/api/projects/{p['id']}/copy").json()
    assert clone["source_project_id"] == p["id"] and clone["scenes"][0]["id"] != p["scenes"][0]["id"]
    assert client.delete(f"/api/projects/{p['id']}").status_code == 200
    assert app.state.service.store.file(clone["id"], "assets/a/test.txt").read_text() == "owned"
    assert client.get(f"/api/projects/{p['id']}").status_code == 404


def test_symbol_and_pause_offsets():
    text = "区间 [0,1]，数组 [1,2,3]，引用 [1]。{{pause:1.5}} E = mc²。"
    spoken, events = speech(text, {"E = mc²": "能量等于质量乘以光速平方"})
    assert all(x in spoken for x in ["[0,1]", "[1,2,3]", "[1]"])
    assert "{{pause:" not in spoken and spoken[events[0]["offset"]:] == " 能量等于质量乘以光速平方。"
    with pytest.raises(Failure):
        speech("bad {{pause:500}}", {})
    with pytest.raises(Failure):
        speech("bad {{pause:abc}}", {})


def test_mapping_reorder_skip_merge_split(app, client):
    p = seed(app)
    ids = [s["id"] for s in p["scenes"]]
    p = cmd(client, p, "structure", {"action": "reorder", "scene_ids": ids[::-1]})
    assert p["scenes"][0]["source_page_ids"] == ["hash:slide-3"]
    p = cmd(client, p, "structure", {"action": "skip", "scene_id": ids[0], "skipped": True})
    assert p["scenes"][-1]["skipped"]
    cmd(client, p, "structure", {"action": "reorder", "scene_ids": [ids[0]] * 3}, 422)
    p = cmd(client, p, "structure", {"action": "merge", "scene_ids": ids[1:][::-1]})
    merged = p["scenes"][0]
    assert merged["source_page_ids"] == ["hash:slide-3", "hash:slide-2"]
    p = cmd(client, p, "structure", {"action": "play_page", "scene_id": merged["id"], "source_page_id": "hash:slide-2"})
    assert p["scenes"][0]["play_page_id"] == "hash:slide-2"
    p = cmd(client, p, "structure", {"action": "split", "scene_id": merged["id"], "parts": ["第一部分", "第二部分"]})
    assert p["scenes"][0]["source_page_ids"] == p["scenes"][1]["source_page_ids"] == merged["source_page_ids"]
    assert p["scenes"][0]["archived_parent"]["merged_from"]


def test_notes_external_ai_missing_and_rollback(app, client):
    p = seed(app)
    sid = p["scenes"][0]["id"]
    p = cmd(client, p, "notes", {"scene_id": sid})
    original = current(p["scenes"][0])
    assert original["display_text"] == p["slides"][0]["notes_original"]
    payload = client.get(f"/api/projects/{p['id']}/import-template").json()
    for item in payload["pages"]:
        item["display_text"] = "外部稿 [1,2,3]，(a+b)，[1]。"
        item["reading_text"] = "外部稿，数组一二三。"
    bad = copy.deepcopy(payload)
    bad["pages"][0]["source_page_id"] = "wrong"
    cmd(client, p, "import", bad, 422)
    assert len(app.state.service.store.get(p["id"])["scenes"][0]["versions"]) == 1
    p = cmd(client, p, "import", payload)
    assert current(p["scenes"][0])["display_text"] == payload["pages"][0]["display_text"]
    p = cmd(client, p, "script", {"scene_id": sid, "display_text": "人工稿"})
    assert p["scenes"][0]["versions"][0] == original
    p = cmd(client, p, "rollback", {"scene_id": sid, "version_id": original["id"]})
    assert current(p["scenes"][0])["display_text"] == original["display_text"]
    assert current(p["scenes"][0])["number"] == 4
    assert client.post(f"/api/projects/{p['id']}/scenes/{sid}/ai", json={"revision": p["revision"]}).status_code == 422


def test_empty_notes_blocked(app, client):
    p = seed(app)
    with app.state.service.store.edit(p["id"]) as p:
        p["slides"][0]["notes_original"] = ""
    p = app.state.service.store.get(p["id"])
    cmd(client, p, "notes", {"scene_id": p["scenes"][0]["id"]}, 422)


def test_review_invalidation_and_snapshot_isolation(app, client):
    p = seed(app)
    first = p["scenes"][0]
    p = cmd(client, p, "script", {"scene_id": first["id"], "display_text": "确认稿"})
    bad = checks(first)
    bad["pending"] = ["公式读法待确认"]
    cmd(client, p, "review", {"scene_id": first["id"], "checks": bad, "note": "教师"}, 422)
    p = cmd(client, p, "review", {"scene_id": first["id"], "checks": checks(first), "note": "测试专用核查"})
    cmd(client, p, "review", {"scene_id": first["id"], "checks": checks(first), "note": "修改锁定稿"}, 409)
    snapshot = app.state.service.store.snapshot(p, {"ready": False, "scope": "isolation_test_only"})
    p = cmd(client, p, "script", {"scene_id": first["id"], "display_text": "编辑中的新稿"})
    assert p["scenes"][0]["confirmed"] is None
    old = client.get(f"/api/projects/{p['id']}/snapshots/{snapshot['snapshot_id']}").json()
    assert current(old["scenes"][0])["display_text"] == "确认稿"
    assert current(old["scenes"][0])["teacher_review"]["note"] == "测试专用核查"
    assert client.post(f"/api/projects/{p['id']}/snapshots", json={"revision": p["revision"]}).status_code == 422


def test_three_page_inheritance_and_media_invalidation(app, client):
    p = seed(app)
    ids = [s["id"] for s in p["scenes"]]
    p = cmd(client, p, "config", {"speed": 1.1, "position": "bottom-left", "terms": {"FFT": "快速傅里叶变换"}})
    for sid in ids:
        p = cmd(client, p, "script", {"scene_id": sid, "display_text": "FFT 与 [0,1]"})
    for sid, vals in zip(ids, [{"show_teacher": False}, {"position": "top-right", "size": .35}, {"speed": .9, "pause_before": 1, "pause_after": 2}]):
        p = cmd(client, p, "page_config", {"scene_ids": [sid], "overrides": vals})
    assert not p["scenes"][0]["effective_config"]["show_teacher"]
    assert p["scenes"][1]["effective_config"]["position"] == "top-right"
    assert p["scenes"][2]["effective_config"]["speed"] == .9
    assert p["scenes"][0]["config_sources"]["speed"] == "project"
    assert current(p["scenes"][0])["display_text"] == "FFT 与 [0,1]"
    assert current(p["scenes"][0])["reading_text"] == "快速傅里叶变换 与 [0,1]"
    p = cmd(client, p, "page_config", {"scene_ids": [ids[2]], "overrides": {}, "replace": True})
    assert p["scenes"][2]["effective_config"]["speed"] == 1.1
    cmd(client, p, "page_config", {"scene_ids": ids, "overrides": {"speed": 99}}, 422)


def test_preflight_missing_input_and_teacher_questions(app, client):
    p = seed(app)
    r = client.post(f"/api/projects/{p['id']}/preflight").json()
    codes = {e["code"] for e in r["blockers"]}
    assert {"INPUT_MISSING", "EMPTY_SCRIPT", "AUDITION_REQUIRED", "SERVICE_CONFIG_MISSING", "TEACHING_CHECK_REQUIRED"} <= codes
    assert not r["ready"] and not r["generation_called"]
    assert r["generation_seconds"] is None and r["generation_fee"] is None
    sid = p["scenes"][0]["id"]
    assert client.post(f"/api/projects/{p['id']}/scenes/{sid}/audition", json={"revision": p["revision"]}).status_code == 422


def test_upload_actual_format_limits_and_path_isolation(app, client):
    p = app.state.service.store.create("上传")
    url = f"/api/projects/{p['id']}/assets/photo"
    data = {"authorization": "工程测试", "teacher_id": "fixture", "revision": p["revision"]}
    assert client.post(url, data=data, files={"file": ("x.png", b"", "image/png")}).status_code == 422
    assert client.post(url, data=data, files={"file": ("x.exe", b"bad", "image/png")}).status_code == 422
    r = client.post(url, data=data, files={"file": ("x.jpg", b"bad", "image/jpeg")})
    assert r.status_code == 200
    svc = app.state.service
    job = svc.store.claim()
    svc.execute(job)
    assert svc.store.jobs(p["id"])[0]["state"] == "failed"
    assert svc.store.get(p["id"])["assets"]["photo"]["state"] == "failed"
    with pytest.raises(Failure):
        svc.store.file(p["id"], "../../../.session-token")


def test_photo_crop_is_real_and_copy_owned(app, client):
    p = app.state.service.store.create("裁剪")
    buf = io.BytesIO()
    Image.new("RGB", (800, 600), (10, 60, 30)).save(buf, format="PNG")
    data = {"authorization": "工程测试图", "teacher_id": "fixture", "revision": p["revision"]}
    assert client.post(f"/api/projects/{p['id']}/assets/photo", data=data, files={"file": ("fixture.png", buf.getvalue())}).status_code == 200
    svc = app.state.service
    svc.execute(svc.store.claim())
    p = svc.store.get(p["id"])
    r = client.post(f"/api/projects/{p['id']}/assets/photo/selection", json={"revision": p["revision"], "selection": {"box": [100, 100, 500, 400]}})
    assert r.status_code == 200, r.text
    svc.execute(svc.store.claim())
    p = svc.store.get(p["id"])
    asset = p["assets"]["photo"]
    assert asset["selected_metadata"]["width"] == 400 and asset["selected_metadata"]["height"] == 300
    assert asset["original"] != asset["selected"]
    assert client.get(f"/api/projects/{p['id']}/files/{asset['selected']}").status_code == 200


def test_job_restart_reports_interruption_and_delete_busy(app, client):
    p = seed(app)
    svc = app.state.service
    svc.store.job(p["id"], "ai", {"scene_id": p["scenes"][0]["id"]})
    assert client.delete(f"/api/projects/{p['id']}").status_code == 409
    claimed = svc.store.claim()
    assert claimed["state"] == "running"
    Store(svc.store.root).recover()
    job = svc.store.jobs(p["id"])[0]
    assert job["state"] == "failed" and job["error"]["code"] == "WORKER_INTERRUPTED"


def test_real_subprocess_crash_persistence(app):
    import subprocess
    import sys
    store = app.state.service.store
    p = store.create("故障进程验证")
    store.job(p["id"], "ai", {"scene_id": "test"})
    store.job(p["id"], "ai", {"scene_id": "test-queued"})
    code = "import sys,os;from mooc_m2.store import Store;s=Store(sys.argv[1]);s.claim();os._exit(7)"
    completed = subprocess.run([sys.executable, "-c", code, str(store.root)], capture_output=True)
    assert completed.returncode == 7
    reopened = Store(store.root)
    reopened.recover()
    jobs = reopened.jobs(p["id"])
    assert {j["state"] for j in jobs} == {"queued", "failed"}
    failed = next(j for j in jobs if j["state"] == "failed")
    assert failed["error"]["code"] == "WORKER_INTERRUPTED"
    assert reopened.claim()["payload"]["scene_id"] == "test-queued"


def test_single_worker_process_lock(app):
    from mooc_m2.service import Service
    first = app.state.service
    second = Service(first.config, first.base)
    first.start()
    try:
        with pytest.raises(Failure, match="已有内容工作进程"):
            second.start()
    finally:
        first.close()


def test_ai_provider_context_and_non_overwrite(app, client, monkeypatch):
    p = seed(app)
    svc = app.state.service
    image = svc.store.file(p["id"], "missing.png")
    Image.new("RGB", (100, 60)).save(image)
    svc.config["ai"] = {"base_url": "http://127.0.0.1:9999/v1", "model": "test-double", "api_key_env": "", "allow_paid": False}
    sid = p["scenes"][0]["id"]
    p = cmd(client, p, "script", {"scene_id": sid, "display_text": "人工版本必须保留"})
    sent = []
    class Response:
        status_code = 200
        @property
        def text(self): return json.dumps(self.json())
        def raise_for_status(self): pass
        def json(self):
            obj = {"display_text": "AI 草稿", "reading_text": "AI 草稿", "knowledge_points": ["知识点"], "questions": [], "terms": [], "pending": ["数值待确认"], "extensions": ["模型补充"]}
            return {"choices": [{"message": {"content": json.dumps(obj)}}]}
    def post(self, url, **kwargs):
        sent.append(kwargs["json"])
        return Response()
    assert client.post(f"/api/projects/{p['id']}/scenes/{sid}/ai", json={"revision": p["revision"]}).status_code == 200
    monkeypatch.setattr("httpx.Client.post", post)
    svc.execute(svc.store.claim())
    edited = svc.store.get(p["id"])
    assert edited["scenes"][0]["versions"][0]["display_text"] == "人工版本必须保留"
    assert current(edited["scenes"][0])["mode"] == "manual"
    assert edited["scenes"][0]["ai_candidate"] == edited["scenes"][0]["versions"][-1]["id"]
    content = sent[0]["messages"][1]["content"]
    assert content[1]["type"] == "image_url"
    assert '相邻页资料：' in content[0]['text']
    assert edited["scenes"][0]["confirmed"] is None


def test_layout_retains_audio_speech_controls_stale_only_affected(app, client):
    p = seed(app)
    for s in p["scenes"]:
        p = cmd(client, p, "script", {"scene_id": s["id"], "display_text": "FFT 测试"})
        p = cmd(client, p, "review", {"scene_id": s["id"], "checks": checks(s), "note": "测试核查"})
    with app.state.service.store.edit(p["id"]) as p:
        for s in p["scenes"]:
            s["audition"] = {"stale": False, "input_key": audition_key(p, s), "duration_seconds": 2}
    p = app.state.service.store.get(p["id"])
    first, second = [s["id"] for s in p["scenes"][:2]]
    p = cmd(client, p, "page_config", {"scene_ids": [first], "overrides": {"position": "top-left", "size": .4}})
    assert not p["scenes"][0]["audition"]["stale"] and p["scenes"][0]["confirmed"]
    p = cmd(client, p, "page_config", {"scene_ids": [first], "overrides": {"speed": 1.4}})
    assert p["scenes"][0]["audition"]["stale"] and p["scenes"][0]["confirmed"]
    assert not p["scenes"][1]["audition"]["stale"] and p["scenes"][1]["confirmed"]
    p = cmd(client, p, "page_config", {"scene_ids": [second], "overrides": {"terms": {"FFT": "快速傅里叶变换"}}})
    assert p["scenes"][1]["confirmed"] is None and p["scenes"][1]["audition"]["stale"]


def test_snapshot_http_success_and_frozen_model_input(app, client, monkeypatch, media):
    from mooc_m1.core import digest, run
    p = seed(app, 1)
    svc = app.state.service
    pid = p["id"]
    photo = svc.store.file(pid, "photo.png")
    Image.new("RGB", (600, 600)).save(photo)
    wav = svc.store.file(pid, "voice.wav")
    run([media.ffmpeg, "-nostdin", "-v", "error", "-f", "lavfi", "-i", "sine=frequency=500:duration=5", wav])
    ppt = svc.store.file(pid, "course.pptx")
    prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(ppt)
    with svc.store.edit(pid) as p:
        p["assets"] = {}
        for role, path in [("pptx", ppt), ("reference_audio", wav), ("photo", photo)]:
            p["assets"][role] = {"id": uid(), "original": path.name, "selected": path.name, "sha256": digest(path), "state": "ready", "authorization": {"confirmed": True}, "teacher_id": "test-fixture", "metadata": {"width": 600, "height": 600}, "selection": {"transcript": "test-only"}}
        s = p["scenes"][0]
        v = add_version(p, s, "manual", "仅用于接口审核测试，不是模型语音证据")
        c = checks(s)
        c["questions"] = [{"question": f"理解测试问题 {i}", "answer": "测试答案", "source_page_id": s["source_page_ids"][0]} for i in range(5)]
        review(s, c, "测试替身审核")
        s["audition"] = {"stale": False, "input_key": audition_key(p, s), "duration_seconds": 5, "path": wav.name, "sha256": digest(wav), "human_review": {"scope": "test_double_only"}}
    monkeypatch.setattr("mooc_m1.adapters.LocalModelAdapter.health", lambda self: {"ready": True, "blockers": [], "quality_verified": False})
    p = svc.store.get(pid)
    r = client.post(f"/api/projects/{pid}/snapshots", json={"revision": p["revision"]})
    assert r.status_code == 200, r.text
    snap = r.json()
    assert snap["schema_version"] == "m2.snapshot.v1" and snap["scenes"][0]["effective_config"]["speed"] == 1
    p = cmd(client, p, "script", {"scene_id": p["scenes"][0]["id"], "display_text": "新的文本"})
    frozen = client.get(f"/api/projects/{pid}/snapshots/{snap['snapshot_id']}").json()
    assert current(frozen["scenes"][0])["display_text"] != "新的文本"


@pytest.mark.parametrize("values", [{"speed": float('nan')}, {"size": True}, {"terms": {"a": ""}}, {"position": "unknown"}, {"invented": 1}])
def test_config_rejects_invalid(values):
    with pytest.raises(Failure):
        config_patch(values)
