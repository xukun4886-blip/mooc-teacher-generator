"""Repeatable real upload/long-course evidence. Never creates teacher approval.

Run prepare once, submit once, collect repeatedly. Existing records are retained.
"""
import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path
import httpx
from pptx import Presentation
from mooc_m1.core import read_json, write_json, digest, stamp
from mooc_m2.store import Store
from mooc_m2.content import fingerprint, current
from mooc_m2.service import Service

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'docs/evidence/M3/long-course-validation.json'


def main():
    global REPORT
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare', 'submit', 'retry', 'collect'])
    parser.add_argument('--mode', choices=['photo', 'video', 'fresh_photo', 'fresh_video'])
    parser.add_argument('--external', action='store_true', help='Collect independent external-script courses')
    parser.add_argument('--report', choices=['fresh-ai'], help='Separate new-upload free-vision smoke evidence')
    parser.add_argument('--recheck', action='store_true', help='Decode and download completed exports again')
    args = parser.parse_args()
    if args.external:
        REPORT = ROOT / 'docs/evidence/M3/external-long-courses.json'
    if args.report:
        if args.external or args.action == 'prepare':
            parser.error('fresh-ai report is collected separately; prepare with m4_fresh_ai_validation')
        REPORT = ROOT / 'docs/evidence/M3/fresh-ai-validation.json'
    store = Store(ROOT / 'storage/m2')
    source_id = read_json(ROOT / 'docs/evidence/M2/formal-lesson-review.json')['project_id']
    source = store.get(source_id)
    with httpx.Client(base_url='http://127.0.0.1:8765', trust_env=False, timeout=180) as client:
        client.post('/api/session', json={'token': store.token}).raise_for_status()
        if args.action == 'prepare':
            folder = ROOT / 'storage/m4-validation'
            folder.mkdir(parents=True, exist_ok=True)
            deck = Presentation(store.file(source_id, source['assets']['pptx']['original']))
            for i in reversed(range(len(deck.slides))):
                if not 4 <= i <= 13:
                    slide_id = deck.slides._sldIdLst[i]
                    deck.part.drop_rel(slide_id.rId)
                    deck.slides._sldIdLst.remove(slide_id)
            ppt = folder / 'lecture2-pages5-14.pptx'
            deck.save(ppt)
            evidence = {'created_at': stamp(), 'scope': 'M3 ten-page real unreviewed drafts, M4 preacceptance only',
                        'source_project_id': source_id, 'source_digest_before': fingerprint(source),
                        'source_pages': list(range(5, 15)), 'ppt_sha256': digest(ppt), 'paths': {},
                        'teacher_confirmation_submitted': False, 'metrics_frozen': False,
                        'authorization': 'Reuse existing user authorization recorded 2026-10-01; one teacher, one course; not full M4 dataset'}
            if REPORT.exists():
                evidence = read_json(REPORT)
            old = read_json(ROOT / 'docs/evidence/M3/one-click-real.json')
            for mode in ['photo', 'video']:
                evidence['paths']['fresh_' + mode] = {'project_id': old['paths'][mode]['project_id'], 'scope': 'existing fresh uploads, two pages; not ten-page baseline'}
                if mode not in evidence['paths']:
                    p = client.post('/api/projects', json={'name': 'M3/M4 十页真实预验收 · ' + mode,
                        'config': {'mode': mode, 'resolution': '1080p', 'target_seconds': 18, 'layout': 'sidebar'}})
                    p.raise_for_status()
                    evidence['paths'][mode] = {'project_id': p.json()['id'], 'scope': '10 original pages, target 180s; actual TTS measured', 'upload_jobs': []}
                pid = evidence['paths'][mode]['project_id']
                write_json(REPORT, evidence)
                for role in [mode, 'pptx', 'reference_audio']:
                    deadline = time.monotonic() + 180
                    while store.busy(pid) and time.monotonic() < deadline:
                        time.sleep(.5)
                    if store.busy(pid):
                        raise RuntimeError('Upload processing did not finish within 180s')
                    if store.get(pid)['assets'].get(role, {}).get('state') == 'ready':
                        continue
                    asset_file = ppt if role == 'pptx' else store.file(source_id, source['assets'][role]['original'])
                    revision = store.get(pid)['revision']
                    with asset_file.open('rb') as stream:
                        response = client.post(f'/api/projects/{pid}/quick-assets/{role}', data={'revision': revision},
                            files={'file': (asset_file.name, stream, 'application/octet-stream')})
                    response.raise_for_status()
                    evidence['paths'][mode]['upload_jobs'].append(response.json())
                    write_json(REPORT, evidence)
                print('uploaded', mode, pid)
        else:
            evidence = read_json(REPORT)
            for mode, entry in evidence['paths'].items():
                if args.mode and args.mode != mode:
                    continue
                pid = entry['project_id']
                if args.action == 'submit' and not mode.startswith('fresh_'):
                    response = client.post(f'/api/projects/{pid}/generate')
                    entry.setdefault('submissions', []).append({'at': stamp(), 'status': response.status_code, 'body': response.json()})
                    response.raise_for_status()
                if args.action == 'retry':
                    state = client.get(f'/api/projects/{pid}/generation').json()
                    if state and state['state'] == 'failed':
                        response = client.post(f"/api/projects/{pid}/generation/{state['id']}/retry")
                        entry.setdefault('retries', []).append({'at': stamp(), 'status': response.status_code, 'body': response.json()})
                state = client.get(f'/api/projects/{pid}/generation')
                state.raise_for_status()
                entry['status'] = state.json()
                p = store.get(pid)
                entry['assets'] = {k: {'state': a['state'], 'sha256': a['sha256'], 'selection':
                    {sk: sv for sk, sv in (a.get('selection') or {}).items() if sk != 'transcript'}} for k, a in p['assets'].items()}
                entry['scripts'] = [{'scene_id': s['id'], 'source_page_ids': s['source_page_ids'], 'version_id': current(s)['id'] if current(s) else None,
                                     'mode': current(s)['mode'] if current(s) else None, 'confirmed': s['confirmed'],
                                     'provider_record': current(s)['provenance'].get('provider_record') if current(s) else None,
                                     'pairing_method': current(s)['provenance'].get('pairing_method') if current(s) else None} for s in p['scenes']]
                entry['course_history'] = [{k: j.get(k) for k in ['id', 'state', 'stage', 'generation', 'created_at', 'started_at', 'finished_at', 'error', 'attempts', 'ai_attempts']} for j in store.jobs(pid) if j['kind'] == 'course']
                if entry['status'] and entry['status'].get('media'):
                    internal = Service(read_json(ROOT / 'config/m2.local.json')).m3.get(pid, entry['status']['media']['id'])
                    entry['media_history'] = {k: internal.get(k) for k in ['generation', 'attempts', 'recoveries', 'recovery_history', 'elapsed_seconds', 'created_at', 'started_at', 'finished_at']}
                    entry['scenes'] = [{k: s.get(k) for k in ['id', 'source_index', 'stage_states', 'cache_hits', 'tts', 'portrait', 'compose']} for s in internal['scenes']]
                    measurements = []
                    for scene in internal['scenes']:
                        for layer in ['tts', 'portrait']:
                            media = scene.get(layer)
                            if not media:
                                continue
                            measurements.append({'scene_id': scene['id'], 'layer': layer, 'provider': media['provider'],
                                'cold_native_elapsed_seconds': media['elapsed_seconds'], 'duration_seconds': media['duration_seconds'],
                                'device_peak_used_mib': (media.get('resources') or {}).get('device_peak_used_mib'),
                                'torch_peak_allocated_mib': (media.get('native_loads') or {}).get('torch_peak_allocated_mib'),
                                'torch_peak_reserved_mib': (media.get('native_loads') or {}).get('torch_peak_reserved_mib'),
                                'native_model_timings': {k:v for k,v in (media.get('native_model') or {}).items() if k.endswith('_seconds')},
                                'cache_hit_on_latest_execution': scene.get('cache_hits', {}).get(layer, {}).get('hit')})
                    entry['performance'] = {'phase': 'each_native_request_is_a_cold_process', 'native_measurements': measurements,
                        'device_peak_used_mib': max((m['device_peak_used_mib'] for m in measurements if m['device_peak_used_mib'] is not None), default=None),
                        'ten_page_same_process_warm_measurement': None,
                        'checkpoint_elapsed_seconds': internal.get('elapsed_seconds'),
                        'elapsed_limit': 'checkpoint total excludes unsaved elapsed interval before abrupt process death',
                        'device_scope': 'entire GPU including desktop, separate from torch allocated/reserved'}
                    course = next(j for j in store.jobs(pid) if j['id'] == entry['status']['id'])
                    if course.get('started_at'):
                        first_generation = course.get('generation', 1) == 1
                        entry['performance']['course_initial_queue_seconds'] = (datetime.fromisoformat(course['started_at']) - datetime.fromisoformat(course['created_at'])).total_seconds() if first_generation else None
                        enqueued = course['created_at'] if first_generation else course['attempts'][-1]['at']
                        entry['performance']['course_latest_generation_queue_seconds'] = (datetime.fromisoformat(course['started_at']) - datetime.fromisoformat(enqueued)).total_seconds()
                        entry['performance']['queue_scope'] = 'Initial start timestamp is overwritten by retries; leave initial queue unknown after retry. Latest queue measured from persisted enqueue attempt.'
                    if internal.get('finished_at') and entry['status']['state'] == 'completed':
                        end = datetime.fromisoformat(internal['finished_at'])
                        entry['performance']['media_job_created_to_finished_wall_seconds'] = (end - datetime.fromisoformat(internal['created_at'])).total_seconds()
                        entry['performance']['course_submission_to_media_finished_wall_seconds'] = (end - datetime.fromisoformat(course['created_at'])).total_seconds()
                    if entry['status']['state'] == 'completed' and (args.recheck or not entry.get('exports')):
                        result = internal['result']
                        checks = {}
                        for kind in ['mp4', 'srt', 'script', 'manifest']:
                            response = client.get(f"/api/projects/{pid}/media-jobs/{internal['id']}/exports/{kind}")
                            response.raise_for_status()
                            assert __import__('hashlib').sha256(response.content).hexdigest() == result['file_hashes'][kind]
                            checks[kind] = {'sha256': result['file_hashes'][kind], 'bytes': len(response.content)}
                        entry['exports'] = checks
                        entry['export_checks_at'] = stamp()
                        entry['media_check'] = Service(read_json(ROOT / 'config/m2.local.json')).media.inspect(store.file(pid, result['path']), 'video')
                        streams = entry['media_check']['streams']
                        ends = {s['codec_type']: float(s.get('start_time', 0)) + float(s['duration']) for s in streams if s['codec_type'] in {'video', 'audio'}}
                        entry['av_end_delta_seconds'] = abs(ends['video'] - ends['audio'])
                        from mooc_m3.timeline import course_timeline, srt
                        pages, cues, duration = course_timeline(internal['scenes'])
                        assert store.file(pid, result['exports']['srt']).read_text(encoding='utf-8') == srt(cues)
                        assert all(0 <= q['start'] < q['end'] <= duration for q in cues)
                        for scene in internal['scenes']:
                            assert scene['portrait']['audio_sha256'] == scene['tts']['sha256']
                            assert scene['tts']['provider'] == 'GPT-SoVITS'
                            assert scene['portrait']['provider'] == ('SadTalker' if internal['payload']['mode'] == 'photo' else 'MuseTalk')
                        assert len({s['portrait']['sha256'] for s in internal['scenes']}) == len(internal['scenes'])
                        entry['current_audio_driving_checks'] = {'all_portrait_audio_hashes_match': True, 'distinct_portrait_count': len(internal['scenes']),
                            'human_lipsync_verified': False}
                        if internal['payload']['mode'] == 'video':
                            feature_hashes = [s['portrait']['native_driving'][0]['whisper_chunks_sha256'] for s in internal['scenes']]
                            assert len(set(feature_hashes)) == len(internal['scenes'])
                            entry['source_sequence_checks'] = {'mode': 'video', 'source_sha256': p['assets']['video']['sha256'],
                                'source_duration_seconds': p['assets']['video']['metadata']['duration_seconds'],
                                'distinct_current_audio_feature_count': len(set(feature_hashes)),
                                'source_kind': 'original video sequence, MuseTalk worker video request', 'human_motion_seams_verified': False}
                        entry['subtitle_integrity'] = {'cue_count': len(cues), 'measured_sentence_timing': True, 'file_equals_manifest_cues': True,
                            'human_timing_checked': False, 'full_display_text_preserved': all(''.join(c['text'] for c in scene['tts']['captions']).replace(' ', '') == scene['version']['display_text'].replace(' ', '').strip() for scene in internal['scenes'])}
                        assert entry['subtitle_integrity']['full_display_text_preserved']
                    if entry['status']['state'] == 'completed':
                        role = internal['payload']['mode']
                        source_hash = p['assets'][role].get('selected_sha256', p['assets'][role]['sha256'])
                        for scene in internal['scenes']:
                            request_path = store.file(pid, scene['portrait']['native_log']).parent / 'worker-request.json'
                            native = read_json(request_path)
                            assert native['role'] == role and role in native['request']
                            assert ('photo' if role == 'video' else 'video') not in native['request']
                            assert digest(native['request'][role]) == source_hash
                            assert digest(native['request']['audio']) == scene['tts']['sha256']
                            if role == 'video':
                                assert scene['portrait']['native_driving'] and all(d['audio_sha256'] == scene['tts']['sha256'] for d in scene['portrait']['native_driving'])
                        entry['actual_native_request_checks'] = {'count': len(internal['scenes']), 'selected_person_source_hashes_match': True,
                            'roles_match_selected_mode': True, 'request_audio_hashes_match': True, 'native_driving_audio_hashes_match': True if role == 'video' else None,
                            'video_requests_do_not_use_photo_path': True if role == 'video' else None, 'human_quality_verified': False}
                print(mode, {'state': entry['status']['state'] if entry['status'] else None, 'scripts': sum(bool(s['version_id']) for s in entry['scripts']),
                    'error': entry['status'].get('error') if entry['status'] else None,
                    'progress': entry['status']['media']['progress'] if entry['status'] and entry['status'].get('media') else None})
            evidence['source_unchanged'] = fingerprint(store.get(source_id)) == evidence['source_digest_before']
            evidence['updated_at'] = stamp()
            write_json(REPORT, evidence)


if __name__ == '__main__':
    main()
