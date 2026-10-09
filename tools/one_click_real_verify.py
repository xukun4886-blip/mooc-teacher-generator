"""Prepare/inspect real short courses. No simulated models or teacher confirmations."""
import argparse
import json
import sys
from pathlib import Path
import httpx
from pptx import Presentation
from pptx.util import Inches, Pt
from mooc_m1.core import read_json, write_json, stamp, digest
from mooc_m2.store import Store

sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).resolve().parents[1]
store=Store(root/'storage/m2')
evidence=root/'docs/evidence/M3/one-click-real.json'
parser=argparse.ArgumentParser();parser.add_argument('action',choices=['prepare','inspect','saved']);args=parser.parse_args()
with httpx.Client(base_url='http://127.0.0.1:8765',trust_env=False,timeout=60) as c:
    c.post('/api/session',json={'token':store.token}).raise_for_status()
    if args.action=='prepare':
        folder=root/'storage/one-click-fixture';folder.mkdir(exist_ok=True)
        ppt=Presentation();ppt.slide_width=Inches(10);ppt.slide_height=Inches(5.625)
        for title,body,notes in [('无线接入','无线接入通过无线信号连接终端与网络。','本页解释无线接入的基本含义：终端通过无线信号接入网络。'),('多址技术','多个终端共享无线资源，需要合理分配。','本页解释多址技术的作用：让多个终端有序共享无线资源，减少相互干扰。')]:
            slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text=title
            slide.placeholders[1].text=body
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for r in para.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(28)
            slide.notes_slide.notes_text_frame.text=notes
        ppt.save(folder/'short-course.pptx')
        source=store.get(read_json(root/'docs/evidence/M2/formal-lesson-review.json')['project_id'])
        value={'created_at':stamp(),'scope':'两条两页短课程，全页显示教师；目标每页8秒，仅用于本轮一键链路验证，不替代正式10页/3分钟评审',
               'source_project_id':source['id'],'source_revision_before':source['revision'], 'source_digest_before':__import__('mooc_m2.content',fromlist=['fingerprint']).fingerprint(source),
               'fixture_ppt':str(folder/'short-course.pptx'),'fixture_ppt_sha256':digest(folder/'short-course.pptx'),
               'materials':{k:str(store.file(source['id'],source['assets'][k]['original'])) for k in ['photo','video','reference_audio']},
               'teacher_confirmation':False,'paths':{}}
        for mode in ['photo','video']:
            response=c.post('/api/projects',json={'name':'一键制作真实验证 · '+('照片' if mode=='photo' else '视频'),
                'config':{'layout':'sidebar','resolution':'1080p','target_seconds':8}});response.raise_for_status()
            value['paths'][mode]={'project_id':response.json()['id']}
        write_json(evidence,value);print(json.dumps(value,ensure_ascii=False,indent=2))
    elif args.action=='saved':
        import copy
        import shutil
        from mooc_m2.content import current, uid
        value=read_json(evidence)
        source=store.get(read_json(root/'docs/evidence/M3/real-draft-media.json')['engineering_project_id'])
        value['saved_script_scope']='现有原页5、7的真实AI摘句讲稿，所有页显示教师，作为继续已有课程的一键媒体验证；新上传自动AI课程因服务429保持失败，不混同证据。'
        for mode in ['photo','video']:
            result=c.post('/api/projects',json={'name':'一键继续已有讲稿 · '+mode,'config':{'layout':'sidebar','resolution':'1080p','mode':mode,'target_seconds':8}});result.raise_for_status()
            pid=result.json()['id'];target=store.folder(pid)
            roles=['pptx','reference_audio',mode]
            assets={r:copy.deepcopy(source['assets'][r]) for r in roles}
            slides=copy.deepcopy(source['slides'])
            files=set()
            for a in assets.values():files.update([a['original'],a['selected']])
            for slide in slides:
                files.update([slide['image'],slide['source_xml']])
            for scene in source['scenes']:
                audition = scene.get('preview_audition')
                if audition and store.file(source['id'], audition['path']).is_file():
                    files.add(audition['path'])
            for relative in files:
                path=store.file(pid,relative);path.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(store.file(source['id'],relative),path)
            scenes=copy.deepcopy(source['scenes'])
            for s in scenes:
                s['id']=uid()
                if not s['skipped'] and not s['overrides'].get('show_teacher',True):s['skipped']=True
                s['overrides']={};s['audition']=None;s['confirmed']=None
                v=current(s)
                if v:v['provenance']['sentence_pairs']=[{'display':v['display_text'],'reading':v['reading_text']}]
            with store.edit(pid) as p:
                p.update(assets=assets,slides=slides,scenes=scenes,engineering_only=True,category='development',source_project_id=source['id'])
            value['paths']['saved_'+mode]={'project_id':pid}
            print('saved_'+mode,pid)
        write_json(evidence,value)
    else:
        value=read_json(evidence)
        for mode,entry in value['paths'].items():
            pid=entry['project_id'];r=c.get(f'/api/projects/{pid}/generation');r.raise_for_status();state=r.json()
            p=store.get(pid)
            entry.update(status=state,asset_states={k:a['state'] for k,a in p['assets'].items()},
                         scripts=[{'source_page_ids':s['source_page_ids'],'confirmed':s['confirmed'],'versions':len(s['versions'])} for s in p['scenes']])
            if state and state.get('media') and state['state']=='completed':
                internal=__import__('mooc_m2.service',fromlist=['Service']).Service(read_json(root/'config/m2.local.json')).m3.get(pid,state['media']['id'])
                entry['real_scenes']=[{k:s.get(k) for k in ['source_index','stage_states','tts','portrait','compose']} for s in internal['scenes']]
                entry['media_attempts'] = internal.get('attempts', [])
                entry['media_generation'] = internal['generation']
                entry['cache_hits'] = [s.get('cache_hits') for s in internal['scenes']]
                entry['media_check']=__import__('mooc_m2.service',fromlist=['Service']).Service(read_json(root/'config/m2.local.json')).media.inspect(store.file(pid,state['media']['result']['path']),'video')
                for kind in ['mp4','srt','script','manifest']:
                    response=c.get(f"/api/projects/{pid}/media-jobs/{state['media']['id']}/exports/{kind}");response.raise_for_status()
                    assert __import__('hashlib').sha256(response.content).hexdigest()==state['media']['result']['file_hashes'][kind]
                entry['four_exports_verified']=True
            print(mode,json.dumps({'assets':entry['asset_states'],'generation':state},ensure_ascii=False))
        source=store.get(value['source_project_id'])
        value['source_unchanged']=__import__('mooc_m2.content',fromlist=['fingerprint']).fingerprint(source)==value['source_digest_before']
        value['updated_at']=stamp();write_json(evidence,value)
