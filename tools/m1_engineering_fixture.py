"""Self-authored engineering fixture, NOT an authorized course or teacher sample.

Real PPTX, native chart, OOXML equation, media poster and notes are sent to the
installed PowerPoint renderer. Test tone/video are solely decoder test stimuli.
"""
import sys
from pathlib import Path
from datetime import datetime, timezone

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

from mooc_m1.core import digest, run, write_json
from mooc_m1.__main__ import setup
from mooc_m1.ppt import parse_pptx, render_powerpoint


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    root = (Path("storage/m1/engineering") / datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%SZ")).resolve()
    root.mkdir(parents=True, exist_ok=True)
    media, _, _ = setup({})
    # A uniform color image is an image-only parsing fixture, not a human portrait.
    poster = root / "image-only.png"
    Image.new("RGB", (800, 450), (37, 68, 98)).save(poster)
    video = root / "decoder-only.mp4"
    run([media.ffmpeg, "-nostdin", "-y", "-f", "lavfi", "-i", "color=c=blue:s=320x180:d=1:r=25",
         "-c:v", "libx264", "-pix_fmt", "yuv420p", video])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.3333), Inches(7.5)
    for i in range(1, 11):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if i == 9:
            s.shapes.add_picture(str(poster), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
            continue
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        box.text = f"M1 工程自建课件 · 原页 {i:02d}"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        body = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11), Inches(2))
        body.text = f"第 {i} 页，原页序号必须保持。\n数字 2026，数组 [1,2,3]，括号 (a+b)，引用 [12]。\n该课件只验证解析和真实原页渲染。"
        for p in body.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Microsoft YaHei"
                r.font.size = Pt(22)
        s.notes_slide.notes_text_frame.text = f"备注原文 {i:02d}：不自动润色，不把 [1,2,3] 删除。"
        if i == 5:
            data = CategoryChartData()
            data.categories = ["甲", "乙", "丙"]
            data.add_series("工程数据", (1, 2, 3))
            s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(8), Inches(3), data)
        if i == 7:
            # Native equation XML, not a pre-rendered replacement picture.
            mathbox = s.shapes.add_textbox(Inches(1), Inches(4), Inches(9), Inches(1))
            p = mathbox.text_frame.paragraphs[0]._p
            wrapper = etree.Element("{http://schemas.microsoft.com/office/drawing/2010/main}m", nsmap={"a14": "http://schemas.microsoft.com/office/drawing/2010/main"})
            mathpara = OxmlElement("m:oMathPara")
            equation = OxmlElement("m:oMath")
            chunk = OxmlElement("m:r")
            properties = OxmlElement("a:rPr")
            properties.set("sz", "3200")
            latin = OxmlElement("a:latin")
            latin.set("typeface", "Cambria Math")
            properties.append(latin)
            chunk.append(properties)
            text = OxmlElement("m:t")
            text.text = "E = mc²"
            chunk.append(text)
            equation.append(chunk)
            mathpara.append(equation)
            wrapper.append(mathpara)
            p.append(wrapper)
        if i == 10:
            s.shapes.add_movie(str(video), Inches(1), Inches(4), Inches(4), Inches(2.25),
                               poster_frame_image=str(poster), mime_type="video/mp4")
    source = root / "engineering-10-pages.pptx"
    prs.save(source)
    parsed = parse_pptx(source, root / "parsed")
    rendered = render_powerpoint(source, root / "rendered")
    assert parsed["page_count"] == rendered["page_count"] == 10
    assert [p["source_index"] for p in parsed["pages"]] == list(range(1, 11))
    for page in parsed["pages"]:
        i = page["source_index"]
        assert page["notes_original"] == ("" if i == 9 else f"备注原文 {i:02d}：不自动润色，不把 [1,2,3] 删除。")
    expected = {5: "CHART_REVIEW_REQUIRED", 7: "FORMULA_REVIEW_REQUIRED",
                9: "IMAGE_ONLY_PAGE", 10: "EMBEDDED_MEDIA_REVIEW_REQUIRED"}
    for page, marker in expected.items():
        assert marker in {a["code"] for a in parsed["pages"][page - 1]["anomalies"]}
    result = {"fixture_kind": "self_authored_engineering_only", "not_teacher_capability_evidence": True,
              "source": {"path": str(source), "sha256": digest(source)},
              "parsed_path": str(root / "parsed/structure.json"), "rendered_path": str(root / "rendered/render.json"),
              "pages": rendered["pages"], "timings": rendered["timings"],
              "checks": {"page_count": 10, "page_order": "matched", "notes_original": "matched",
                         "native_chart_flag": True, "native_equation_flag": True,
                         "image_only_flag": True, "embedded_media_flag": True},
              "renderer": rendered["renderer"], "renderer_version": rendered["renderer_version"],
              "authorized_course_review": "pending", "stage_passed": False}
    write_json("docs/evidence/M1/ppt-engineering.json", result)
    print("Real PowerPoint rendered 10 engineering pages; M1 course validation remains pending.")


if __name__ == "__main__":
    main()
